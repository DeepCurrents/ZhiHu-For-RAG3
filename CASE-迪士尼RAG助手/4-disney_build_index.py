# -*- coding: utf-8 -*-
"""
迪士尼RAG助手 - 知识库构建（索引入库）

【这是整个 RAG 的"离线阶段"】
RAG 分两个阶段，务必在脑子里把这条线划清楚：

    离线（本文件）：原始资料 → 解析 → 切片 → 向量化 → 建索引 → 落盘
    在线（5-disney_query.py）：用户问题 → 向量化 → 检索索引 → 拼上下文 → LLM 生成答案

离线阶段可以慢慢跑、可以重跑；在线阶段要求快。两者唯一的强耦合是：
**必须使用同一个 Embedding 模型**，否则查询向量和库内向量不在同一坐标系，检索结果就是噪声。

【本文件要解决的三个问题】
1. 异构数据统一化：知识库里同时有 Word 文档、图片、视频，格式完全不同，
   怎么让它们"检索时能放在一起比"？——统一到同一个文本向量空间：
       .docx → 解析文本 → 切片        → 文本向量
       图片  → 视觉模型转文字描述      → 文本向量
       视频  → 视觉模型转文字描述      → 文本向量
   这就是本案例"多模态 RAG"的核心技巧：把多模态问题降维成文本问题。

2. 资料与向量的对应关系：向量库只存向量，不知道它来自哪个文件。
   所以额外维护一份 metadata（元数据）列表，下标与向量一一对应：
       all_vectors[i]  ←→  metadata_store[i]
   检索命中第 i 条，就能回头查出它的来源、类型、原文。

3. 持久化：FAISS 索引和元数据分别存成两个文件，查询时再加载，避免每次重复调用模型。

【运行前提】
- 环境变量 AGICTO_API_KEY
- 依赖：faiss_cpu、numpy、openai、python_docx
- 知识库目录 disney_knowledge_base/ 下放 .docx，images/ 子目录下放图片

【注意】只要换了 Embedding 模型，就必须重跑本文件重建索引。
"""

import os
import base64
import json
import numpy as np
import faiss
from openai import OpenAI
from docx import Document as DocxDocument

# agicto 配置（OpenAI 兼容接口）
AGICTO_BASE_URL = "https://api.agicto.cn/v1"
# 可按 agicto 账户实际可用列表替换
AGICTO_VISION_MODEL = "gemini-2.5-pro"
AGICTO_EMBEDDING_MODEL = "text-embedding-3-small"

# 当前脚本所在目录，保证从任意工作目录运行都能定位到资源文件
# （__file__ 在 Jupyter 里不存在，所以用 globals() 判断后回退到当前工作目录）
BASE_DIR = os.path.dirname(os.path.abspath(__file__)) if "__file__" in globals() else os.getcwd()

# 建库是耗时操作，Key 没配好就没必要往下跑，所以在入口处直接"快速失败"
AGICTO_API_KEY = os.getenv("AGICTO_API_KEY")
if not AGICTO_API_KEY:
    raise ValueError("错误：请设置 'AGICTO_API_KEY' 环境变量。")

client = OpenAI(
    api_key=AGICTO_API_KEY,
    base_url=AGICTO_BASE_URL,
)

# 输入资源目录
DOCS_DIR = os.path.join(BASE_DIR, "disney_knowledge_base")
IMG_DIR = os.path.join(DOCS_DIR, "images")

# 输出文件
# 索引存向量、元数据存"向量对应的原文与来源"，两者必须配对使用、同版本
INDEX_FILE = os.path.join(BASE_DIR, "disney_index.faiss")
METADATA_FILE = os.path.join(BASE_DIR, "disney_metadata.json")

# 切分参数
CHUNK_SIZE = 500  # 每个chunk的字符数
CHUNK_OVERLAP = 50  # chunk之间的重叠字符数

# 视频知识库
# 视频无法像文档那样直接遍历目录（要么在公网、要么体积太大），
# 所以用一个手工维护的清单登记：url + 人类可读的描述
VIDEO_KNOWLEDGE = [
    {
        "url": "https://dataset-1255932437.cos.ap-nanjing.myqcloud.com/mp4/car.mp4",
        "description": "汽车剐蹭视频"
    }
]


def parse_docx(file_path):
    """解析 DOCX 文件，提取全部文本

    DOCX 是 zip + XML 结构，python-docx 提供了高层 API，但段落和表格需要分开处理。
    这里直接遍历 body 里的 XML 元素：<w:p> 是段落，<w:tbl> 是表格。

    为什么要单独处理表格？因为表格里的信息（价格表、时刻表）往往是问答的高频目标，
    如果只取段落文本就会整块丢失。这里把表格转成 Markdown 表格，
    既保留了行列结构，也让后续切片和 LLM 阅读更自然。
    """
    doc = DocxDocument(file_path)
    all_text = []

    # 按文档中元素的真实顺序遍历，保证"表格出现在段落之间"的位置关系不丢
    for element in doc.element.body:
        if element.tag.endswith('p'):
            # 段落文本被拆成多个 run（同一段落里不同格式的片段），需要拼接
            paragraph_text = ""
            for run in element.findall('.//w:t', {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):
                paragraph_text += run.text if run.text else ""
            if paragraph_text.strip():
                all_text.append(paragraph_text.strip())

        elif element.tag.endswith('tbl'):
            # doc.tables 是按顺序排列的，用 _element 反查当前表格对象
            table = [t for t in doc.tables if t._element is element][0]
            if table.rows:
                # 转成 Markdown 表格：第一行是表头，第二行是分隔线
                md_table = []
                header = [cell.text.strip() for cell in table.rows[0].cells]
                md_table.append("| " + " | ".join(header) + " |")
                md_table.append("|" + "---|"*len(header))
                for row in table.rows[1:]:
                    row_data = [cell.text.strip() for cell in row.cells]
                    md_table.append("| " + " | ".join(row_data) + " |")
                all_text.append("\n".join(md_table))

    return "\n".join(all_text)


def split_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """按固定长度切分文本

    最朴素的切片：不认标点、不做句子对齐，纯按位置切。
    之所以在这里够用，是因为文档本身段落清晰、每段都不长，
    配合 50 字符的重叠已能避免关键信息被切断。
    （想深入对比各种切片策略，见 CASE-切片策略 目录）
    """
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        # 跳过空白块，避免产生无意义的向量
        if chunk.strip():
            chunks.append(chunk.strip())
        # 下一块从「本块结束位置 - overlap」开始，实现重叠
        start = end - overlap
    return chunks


def describe_media(content_list):
    """用视觉模型把图片/视频转成文字描述
    agicto 的 /v1/embeddings 只接受文本输入，没有多模态向量化接口
    """
    resp = client.chat.completions.create(
        model=AGICTO_VISION_MODEL,
        messages=[{"role": "user", "content": content_list}],
    )
    return resp.choices[0].message.content


def get_text_embedding(text):
    """文本embedding"""
    resp = client.embeddings.create(
        model=AGICTO_EMBEDDING_MODEL,
        input=text,
    )
    return resp.data[0].embedding


def get_image_embedding(image_path):
    """图片embedding：先转文字描述，再对描述做文本向量化"""
    # 读二进制 → base64，把图片内嵌进请求体，无需图床
    with open(image_path, "rb") as f:
        base64_image = base64.b64encode(f.read()).decode('utf-8')

    # 从扩展名推断 MIME 类型：jpg 的标准写法是 jpeg，需要特判；
    # 其它格式（png/bmp/gif）直接转小写即可
    ext = os.path.splitext(image_path)[1].lower().lstrip('.')
    if ext == 'jpg':
        ext = 'jpeg'
    image_data = f"data:image/{ext};base64,{base64_image}"

    description = describe_media([
        {"type": "text", "text": "请详细描述这张图片，用于后续检索：包含画面主体、场景、人物、颜色和图中出现的文字。"},
        {"type": "image_url", "image_url": {"url": image_data}},
    ])
    # 返回描述本身也要留着：写进 metadata，方便出问题时排查"是不是描述写偏了"
    return description, get_text_embedding(description)


def get_video_embedding(video_url):
    """视频embedding：先转文字描述，再对描述做文本向量化"""
    description = describe_media([
        {"type": "text", "text": "请详细描述这个视频，用于后续检索：包含画面主体、场景、发生的事件和关键对话。"},
        # 视频链接同样放在 image_url.url 字段（服务端会自己拉取并解析）
        {"type": "image_url", "image_url": {"url": video_url}},
    ])
    return description, get_text_embedding(description)


def build_and_save():
    """构建知识库并保存

    整体流程：三类资料各自向量化 → 汇总成一个向量矩阵 → 建 FAISS 索引 → 落盘
    关键约束：向量列表与元数据列表必须"同序、等长、一一对应"
    """
    print("\n--- 构建多模态知识库 ---")
    print(f"切分参数: chunk_size={CHUNK_SIZE}, overlap={CHUNK_OVERLAP}")

    # 这两个列表是本函数的核心：同一下标代表同一条知识
    metadata_store = []
    all_vectors = []
    # 全局自增 ID，跨文档、图片、视频统一编号，保证 metadata["id"] 唯一
    doc_id = 0

    # 处理Word文档
    for filename in os.listdir(DOCS_DIR):
        # 跳过隐藏文件和子目录（images/ 是子目录，由下面单独处理）
        if filename.startswith('.') or os.path.isdir(os.path.join(DOCS_DIR, filename)):
            continue

        file_path = os.path.join(DOCS_DIR, filename)
        if filename.endswith(".docx"):
            print(f"  处理文档: {filename}")
            full_text = parse_docx(file_path)
            # 注意顺序：先解析出全文，再切片，最后逐块向量化
            chunks = split_text(full_text)
            print(f"    文档长度: {len(full_text)} 字符, 切分为 {len(chunks)} 个chunk")

            for chunk in chunks:
                # 一个 chunk 对应一条知识：既存向量，也存原文和来源
                metadata = {
                    "id": doc_id,
                    "source": filename,
                    "type": "text",
                    "content": chunk
                }

                # 这里是一次网络调用。真实项目应改成批量 embedding（input 传数组）+ 失败重试
                vector = get_text_embedding(chunk)
                all_vectors.append(vector)
                metadata_store.append(metadata)
                doc_id += 1

    # 处理图片（不再需要OCR，视觉模型生成的描述已包含图片语义）
    print("  处理图片...")
    for img_filename in os.listdir(IMG_DIR):
        # 过滤出常见图片格式，忽略目录里的其它杂项文件
        if img_filename.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp')):
            img_path = os.path.join(IMG_DIR, img_filename)
            print(f"    - {img_filename}")

            description, vector = get_image_embedding(img_path)

            metadata = {
                "id": doc_id,
                "source": f"图片: {img_filename}",
                "type": "image",
                # path 留着是为了答应用户"想看原图"时能定位到文件
                "path": img_path,
                "description": description,
                # content 存的是简短标识而不是描述全文，避免上下文过长
                "content": f"[图片] {img_filename}"
            }

            all_vectors.append(vector)
            metadata_store.append(metadata)
            doc_id += 1

    # 处理视频
    print("  处理视频...")
    for video_info in VIDEO_KNOWLEDGE:
        print(f"    - {video_info['description']}")

        description, vector = get_video_embedding(video_info["url"])

        metadata = {
            "id": doc_id,
            "source": f"视频: {video_info['description']}",
            "type": "video",
            "url": video_info["url"],
            "description": description,
            "content": f"[视频] {video_info['description']}"
        }

        all_vectors.append(vector)
        metadata_store.append(metadata)
        doc_id += 1

    # 创建FAISS索引
    if all_vectors:
        # 维度从模型返回的向量长度取得，不写死数字——换 Embedding 模型时这里自动适配
        dim = len(all_vectors[0])
        print(f"\n向量维度: {dim}")

        # IndexFlatL2：最基础的"暴力检索"索引，逐个算 L2（欧氏）距离取最近邻
        # 优点是精确、无需训练；缺点是数据量大时慢。数据上百万级才需要考虑 IVF/HNSW 等近似索引
        index = faiss.IndexFlatL2(dim)
        # FAISS 只接受 float32，所以必须显式 astype；形状是 (n, dim)
        index.add(np.array(all_vectors).astype('float32'))

        # 保存索引
        faiss.write_index(index, INDEX_FILE)
        print(f"索引已保存: {INDEX_FILE}")

        # 元数据是普通 JSON，和索引分开存。查询时两者一起加载，靠下标对齐
        with open(METADATA_FILE, 'w', encoding='utf-8') as f:
            json.dump(metadata_store, f, ensure_ascii=False, indent=2)
        print(f"元数据已保存: {METADATA_FILE}")

    # 统计
    # 三类资料分别计数，方便确认"图片/视频是不是漏处理了"（比如路径写错导致 0 条）
    text_count = sum(1 for m in metadata_store if m["type"] == "text")
    image_count = sum(1 for m in metadata_store if m["type"] == "image")
    video_count = sum(1 for m in metadata_store if m["type"] == "video")
    print(f"\n完成! 文本:{text_count}, 图片:{image_count}, 视频:{video_count}")

    # 打印所有知识条目
    print("\n--- 知识库内容 ---")
    for m in metadata_store:
        # :2d 让 ID 右对齐、:5s 让类型左对齐，输出整齐便于肉眼检查
        print(f"\n[{m['id']:2d}] [{m['type']:5s}] {m.get('source', '')}")
        print(f"    {m['content']}")


if __name__ == "__main__":
    build_and_save()
